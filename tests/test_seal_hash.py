"""Bug 3: the seal fingerprint is computed once, at the moment of sealing, persisted, and reused by every export."""
import hashlib
import io
import json
import time
import zipfile

import docx

import database
import drafts_service as svc
import export_engine

SECTIONS = {"sections": [
    {"number": 1, "title": "Facts", "content": "The petitioner approached the court.\n\nSecond paragraph."},
    {"number": 2, "title": "Order", "content": "Petition allowed."},
]}


def _save(client, h, payload, **extra):
    r = client.post("/api/save_draft", json={**payload, **extra}, headers=h)
    assert r.status_code == 200, r.text
    return r.json()["draft_id"]


def _status(client, h, did, status):
    return client.post("/api/draft/update_status", json={"draft_id": did, "status": status}, headers=h)


def _sealed_draft(client, hdr, draft_payload, title="Seal v. State"):
    did = _save(client, hdr["Clerk"], {**draft_payload, "case_title": title,
                                      "draft_content_json": SECTIONS})
    for st in ("IN_REVIEW", "BENCH_APPROVED", "FINAL_SEALED"):
        who = "Clerk" if st == "IN_REVIEW" else "Justice"
        assert _status(client, hdr[who], did, st).status_code == 200
    return did


def _row(did):
    conn = database.get_connection()
    try:
        return dict(conn.execute("SELECT * FROM draft_judgments WHERE id = ?", (did,)).fetchone())
    finally:
        conn.close()


def _docx(client, h, did):
    r = client.get(f"/api/draft/export/docx/{did}", headers=h)
    assert r.status_code == 200
    return r


def test_seal_persists_hash_and_timestamp(client, hdr, draft_payload):
    did = _sealed_draft(client, hdr, draft_payload)
    row = _row(did)
    assert row["sealed_hash"] and len(row["sealed_hash"]) == 64
    assert row["sealed_at"] and row["sealed_hash_backfilled"] == 0
    assert row["sealed_hash"] == export_engine.compute_seal_hash(row)
    api = client.get(f"/api/draft/{did}", headers=hdr["Clerk"]).json()
    assert api["sealed_hash"] == row["sealed_hash"] and api["sealed_hash_backfilled"] is False
    listed = {d["id"]: d for d in client.get("/api/drafts", headers=hdr["Clerk"]).json()}[did]
    assert listed["sealed_hash"] == row["sealed_hash"] and listed["sealed_hash_backfilled"] is False


def test_unsealed_draft_has_no_seal_columns(client, hdr, draft_payload):
    did = _save(client, hdr["Clerk"], draft_payload)
    row = _row(did)
    assert row["sealed_hash"] is None and row["sealed_at"] is None


def test_sealed_docx_is_byte_identical_across_time(client, hdr, draft_payload):
    did = _sealed_draft(client, hdr, draft_payload)
    a = _docx(client, hdr["Clerk"], did)
    time.sleep(2.1)
    b = _docx(client, hdr["Clerk"], did)
    assert hashlib.sha256(a.content).hexdigest() == hashlib.sha256(b.content).hexdigest()
    assert a.content == b.content
    assert a.headers["x-gyanam-sealed"] == "true"
    assert a.headers["x-gyanam-sealed-backfilled"] == "false"


def test_docx_and_print_view_report_same_hash(client, hdr, draft_payload):
    did = _sealed_draft(client, hdr, draft_payload)
    sealed_hash = _row(did)["sealed_hash"]
    d = _docx(client, hdr["Clerk"], did)
    p = client.get(f"/api/draft/export/print/{did}", headers=hdr["Clerk"])
    assert d.headers["x-gyanam-sealed-hash"] == p.headers["x-gyanam-sealed-hash"] == sealed_hash
    assert p.headers["x-gyanam-sealed"] == "true"
    assert sealed_hash in p.text
    text = "\n".join(par.text for t in docx.Document(io.BytesIO(d.content)).tables
                     for row in t.rows for c in row.cells for par in c.paragraphs)
    assert sealed_hash in text


def test_sealed_dates_come_from_sealed_at_not_now(client, hdr, draft_payload):
    did = _sealed_draft(client, hdr, draft_payload)
    conn = database.get_connection()
    conn.execute("UPDATE draft_judgments SET sealed_at = ? WHERE id = ?", ("2024-03-05T10:00:00+00:00", did))
    conn.commit()
    conn.close()
    p = client.get(f"/api/draft/export/print/{did}", headers=hdr["Clerk"]).text
    assert "Pronounced on 05 March 2024" in p
    body = "\n".join(par.text for par in docx.Document(io.BytesIO(_docx(client, hdr["Clerk"], did).content)).paragraphs)
    assert "05 MARCH 2024" in body and "05/03/2024" in body


def test_blocked_edit_and_reseal_never_touch_hash(client, hdr, draft_payload):
    did = _sealed_draft(client, hdr, draft_payload)
    before = _row(did)
    r = client.post("/api/save_draft", json={**draft_payload, "draft_id": did,
                    "draft_content_json": {"sections": [{"number": 1, "title": "X", "content": "tampered"}]}},
                    headers=hdr["Clerk"])
    assert r.status_code == 409
    _status(client, hdr["Justice"], did, "FINAL_SEALED")  # FINAL_SEALED -> FINAL_SEALED via the HTTP path
    svc.update_status_with_reason(did, "FINAL_SEALED", None, None, {"full_name": "J"})  # and via the service
    after = _row(did)
    for k in ("sealed_hash", "sealed_at", "sealed_hash_backfilled", "draft_content_json"):
        assert before[k] == after[k]


def test_seal_via_service_path_persists(client, hdr, draft_payload):
    did = _save(client, hdr["Clerk"], {**draft_payload, "draft_content_json": SECTIONS})
    for st in ("IN_REVIEW", "BENCH_APPROVED", "FINAL_SEALED"):
        svc.update_status_with_reason(did, st, None, None, {"full_name": "Justice"})
    row = _row(did)
    assert row["sealed_hash"] == export_engine.compute_seal_hash(row) and row["sealed_at"]


def test_preview_export_is_marked_unsealed_and_tracks_edits(client, hdr, draft_payload):
    did = _save(client, hdr["Clerk"], {**draft_payload, "draft_content_json": SECTIONS})
    p1 = client.get(f"/api/draft/export/print/{did}", headers=hdr["Clerk"])
    d1 = _docx(client, hdr["Clerk"], did)
    assert p1.headers["x-gyanam-sealed"] == "false" and d1.headers["x-gyanam-sealed"] == "false"
    assert "x-gyanam-sealed-hash" not in p1.headers
    h1 = export_engine.seal_info(svc.get_draft(did))["hash"]
    assert h1 in p1.text
    edited = {"sections": [{"number": 1, "title": "Facts", "content": "Changed text."}]}
    _save(client, hdr["Clerk"], {**draft_payload, "draft_content_json": edited}, draft_id=did, version=1)
    h2 = export_engine.seal_info(svc.get_draft(did))["hash"]
    assert h1 != h2
    p2 = client.get(f"/api/draft/export/print/{did}", headers=hdr["Clerk"])
    assert h2 in p2.text and h1 not in p2.text


def test_backfill_migration(client, hdr, draft_payload):
    conn = database.get_connection()
    cur = conn.execute(
        "INSERT INTO draft_judgments (case_title, case_number, court, date_created, petitioner, respondent, "
        "draft_content_json, status, version) VALUES ('Old v. State','OLD/1','HC','2023-01-02 09:30','Old','State',?, 'SIGNED', 3)",
        (json.dumps(SECTIONS),))
    did = cur.lastrowid
    conn.commit()
    conn.execute("DELETE FROM schema_migrations WHERE name = '011_backfill_seal_hash'")
    conn.commit()
    assert database.run_migrations(conn) == ["011_backfill_seal_hash"]
    row = dict(conn.execute("SELECT * FROM draft_judgments WHERE id = ?", (did,)).fetchone())
    assert row["sealed_hash"] == export_engine.compute_seal_hash(row)
    assert row["sealed_hash_backfilled"] == 1 and row["sealed_at"].startswith("2023-01-02T09:30")
    # idempotent: forcing a second run changes nothing
    conn.execute("DELETE FROM schema_migrations WHERE name = '011_backfill_seal_hash'")
    conn.commit()
    database.run_migrations(conn)
    assert dict(conn.execute("SELECT * FROM draft_judgments WHERE id = ?", (did,)).fetchone()) == row
    conn.close()
    api = client.get(f"/api/draft/{did}", headers=hdr["Clerk"]).json()
    assert api["sealed_hash_backfilled"] is True and api["sealed_hash"] == row["sealed_hash"]
    d = _docx(client, hdr["Clerk"], did)
    assert d.headers["x-gyanam-sealed"] == "true" and d.headers["x-gyanam-sealed-backfilled"] == "true"
    assert d.headers["x-gyanam-sealed-hash"] == row["sealed_hash"]


def test_sealed_without_persisted_hash_agrees_across_paths(client, hdr, draft_payload):
    conn = database.get_connection()
    did = conn.execute(
        "INSERT INTO draft_judgments (case_title, case_number, court, date_created, petitioner, respondent, "
        "draft_content_json, status, version) VALUES ('Late v. State','LATE/1','HC','2022-05-06 10:00','L','S',?, 'FINAL_SEALED', 1)",
        (json.dumps(SECTIONS),)).lastrowid
    conn.commit()
    conn.close()
    d = _docx(client, hdr["Clerk"], did)
    p = client.get(f"/api/draft/export/print/{did}", headers=hdr["Clerk"])
    assert d.headers["x-gyanam-sealed-hash"] == p.headers["x-gyanam-sealed-hash"]
    assert d.headers["x-gyanam-sealed-backfilled"] == "true"


def test_docx_is_valid_zip_with_pinned_timestamps(client, hdr, draft_payload):
    did = _sealed_draft(client, hdr, draft_payload)
    raw = _docx(client, hdr["Clerk"], did).content
    z = zipfile.ZipFile(io.BytesIO(raw))
    assert z.testzip() is None
    assert {i.date_time for i in z.infolist()} == {(1980, 1, 1, 0, 0, 0)}
    d = docx.Document(io.BytesIO(raw))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "The petitioner approached the court." in text and "Petition allowed." in text
    cp = d.core_properties
    assert cp.created == cp.modified


def test_normalize_docx_zip_preserves_order_and_compression():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("b.txt", "bbb" * 50)
        z.writestr("a.txt", "aaa")
    out = zipfile.ZipFile(io.BytesIO(export_engine.normalize_docx_zip(buf.getvalue())))
    assert [i.filename for i in out.infolist()] == ["b.txt", "a.txt"]
    assert all(i.compress_type == zipfile.ZIP_DEFLATED for i in out.infolist())
    assert out.read("b.txt") == b"bbb" * 50
