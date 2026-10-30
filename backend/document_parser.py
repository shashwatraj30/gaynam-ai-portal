"""Text extraction and honest field parsing for uploaded pleadings.

Nothing here invents content: a field that cannot be found in the text is returned as '' (or [] for issues).
"""
import os
import re
from typing import Any, Dict, List, Tuple

from pypdf import PdfReader

FIELD_KEYS = ("court", "case_number", "case_title", "petitioner", "respondent", "facts", "prayer",
              "statutory_provisions", "petitioner_arguments", "respondent_arguments")


def extract_text_from_file(file_path: str) -> str:
    """Raw text of a PDF / DOCX / image / text file. Returns '' when nothing could be read (never placeholder text)."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        try:
            reader = PdfReader(file_path)
            joined = "\n".join(t for t in (p.extract_text() for p in reader.pages) if t).strip()
            if len(joined) >= 50:
                return joined
        except Exception:
            joined = ""
        try:
            from ocr_engine import process_scanned_brief
            res = process_scanned_brief(file_path)
            return (res or {}).get("full_text") or joined
        except Exception:
            return joined
    if ext in (".docx", ".doc"):
        try:
            import docx
            doc = docx.Document(file_path)
            parts = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_txt:
                        parts.append(row_txt)
            return "\n".join(parts)
        except Exception:
            return ""
    if ext in (".jpg", ".jpeg", ".png", ".webp", ".tiff", ".tif", ".bmp"):
        try:
            from ocr_engine import process_scanned_brief
            res = process_scanned_brief(file_path)
            return (res or {}).get("full_text") or ""
        except Exception:
            return ""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return ""


# ---------------------------------------------------------------- field extraction
_COURT_RE = re.compile(r"\b(HIGH COURT|SUPREME COURT|DISTRICT COURT|SESSIONS COURT|TRIBUNAL|COURT OF)\b", re.I)
_CASE_NO_RE = re.compile(
    r"\b(?:(?:S\.?\s?B\.?|D\.?\s?B\.?)\s+)?(?:[A-Za-z.]+\s+){0,4}"
    r"(?:PETITION|APPEAL|SUIT|APPLICATION|REVISION|REFERENCE|CASE|W\.?P\.?|SLP|CRLMP)\b[^\n]{0,40}?"
    r"\bNO\.?\s*[:.]?\s*\d[\w./\- ]*?(?:\s+OF\s+\d{4}|/\d{2,4})?(?=\s*$|\s*[,;(])", re.I)
_VERSUS_LINE_RE = re.compile(r"^\s*(?:versus|vs\.?|v\.|v/s)\s*$", re.I)
_VERSUS_INLINE_RE = re.compile(r"^(.{3,160}?)\s+(?:versus|vs\.?|v\.|v/s)\s+(.{3,200})$", re.I)
_PARTY_TAIL_RE = re.compile(r"\s*\.{2,}\s*(?:the\s+)?(?:PETITIONERS?|APPELLANTS?|RESPONDENTS?|DEFENDANTS?|PLAINTIFFS?|APPLICANTS?|OPPOSITE PARTY).*$", re.I)
_PARTY_ROLE_RE = re.compile(r"\b(PETITIONERS?|APPELLANTS?|RESPONDENTS?|DEFENDANTS?|PLAINTIFFS?|APPLICANTS?)\b\s*\.?\s*$", re.I)
_PERSON_TAIL_RE = re.compile(r",?\s+(?:S|D|W|R)/o\b.*$|,\s*(?:resident|r/o)\b.*$", re.I)

_HEADINGS = (
    ("FACTS", re.compile(r"^(?:[A-Z]\.\s*)?(?:synopsis(?:\s+and\s+facts)?|(?:brief\s+)?facts(?:\s+of\s+the\s+case)?(?:\s*(?:&|and)\s*grounds)?|"
                         r"substance\s+of\s+(?:the\s+)?complaint(?:\s*/\s*police\s+action)?|statement\s+of\s+facts)\s*:?\s*$", re.I)),
    ("ISSUES", re.compile(r"^(?:[A-Z]\.\s*)?(?:questions?\s+of\s+law|issues?(?:\s+for\s+determination)?|points?\s+for\s+determination|"
                          r"grounds\s+of\s+challenge)\s*:?\s*$", re.I)),
    ("ARGS_P", re.compile(r"^(?:[A-Z]\.\s*)?(?:(?:submissions?|arguments?|contentions?)\s+(?:on\s+behalf\s+of|of)\s+(?:the\s+)?"
                          r"(?:petitioners?|applicants?|appellants?)|(?:petitioners?|applicants?|appellants?)['’]?s?\s+"
                          r"(?:submissions?|arguments?|contentions?))\s*:?\s*$", re.I)),
    ("ARGS_R", re.compile(r"^(?:[A-Z]\.\s*)?(?:(?:submissions?|arguments?|contentions?)\s+(?:on\s+behalf\s+of|of)\s+(?:the\s+)?"
                          r"(?:respondents?|opposite\s+party|state)|(?:respondents?|state)['’]?s?\s+"
                          r"(?:submissions?|arguments?|contentions?))\s*:?\s*$", re.I)),
    ("GROUNDS", re.compile(r"^(?:[A-Z]\.\s*)?grounds?\s*:?\s*$", re.I)),
    ("PRAYER", re.compile(r"^(?:[A-Z]\.\s*)?(?:prayer|relief\s+sought|reliefs?)\s*:?\s*$", re.I)),
    ("END", re.compile(r"^(?:investigation\s+directed|verification|affidavit|annexures?|list\s+of\s+(?:dates|documents)|"
                       r"(?:acts?\s*&\s*sections|seized\s+articles|accused\s+(?:persons|particulars)|complainant\s*/\s*informant|"
                       r"general\s+diary[^:]*)\s*:?)\s*:?\s*$", re.I)),
)
_QUESTION_RE = re.compile(r"\bWhether\b[^?]{10,500}\?", re.I)
_STATUTE_RE = re.compile(
    r"(?i:sections?|articles?|rules?|orders?|u/s\.?)\s+\d+[A-Za-z]{0,2}(?:\([\w]+\))*"
    r"(?:\s*(?:,|&|and|/)\s*\d+[A-Za-z]{0,2}(?:\([\w]+\))*)*"
    r"(?:\s+(?:(?i:of\s+(?:the\s+)?)(?:(?i:constitution of india)|(?i:constitution)|(?:[A-Z][A-Za-z&'.]*\s+){1,6}?(?:Act|Code|Rules|Sanhita)(?:,?\s*\d{4})?)"
    r"|Cr\.P\.C\.|C\.P\.C\.|I\.P\.C\.|N\.I\. Act|BNSS|BNS|CrPC|CPC|IPC))?")


def _clean_line(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def _heading_kind(line: str) -> str:
    l = line.strip()
    if not l or len(l) > 70:
        return ""
    for kind, rx in _HEADINGS:
        if rx.match(l):
            return kind
    return ""


def _sections(lines: List[str]) -> Dict[str, List[str]]:
    out: Dict[str, List[str]] = {}
    cur = ""
    for line in lines:
        kind = _heading_kind(line)
        if kind:
            cur = "" if kind == "END" else kind
            if cur:
                out.setdefault(cur, [])
            continue
        if cur:
            out[cur].append(line)
    return {k: v for k, v in out.items() if any(x.strip() for x in v)}


# FIR-form heading anchors. Hoisted as raw fragments so _parse_fir() and the classifier compile from the SAME
# source text and cannot drift apart.
_PS_HEAD = r"POLICE STATION\s*:"
_ACCUSED_WORDS = r"ACCUSED\s+(?:PERSONS|PARTICULARS)"
_ACCUSED_HEAD = _ACCUSED_WORDS + r"\s*:"
_ACTS_HEAD = r"ACTS\s*&\s*SECTIONS\s*:?"
_SUBSTANCE_HEAD = r"SUBSTANCE OF COMPLAINT"
_FIR_POLICE_STATION_RE = re.compile(_PS_HEAD + r"\s*([^\n]+)", re.I)

# Classifier markers: headings are line-anchored so FIR boilerplate quoted inside a pleading's argument text
# does not count; the two phrase markers may appear anywhere.
_FIR_FORM_MARKERS: Tuple[re.Pattern, ...] = (
    re.compile(r"^\s*" + _PS_HEAD, re.I | re.M),
    re.compile(r"^\s*" + _ACCUSED_HEAD, re.I | re.M),
    re.compile(r"^\s*" + _ACTS_HEAD, re.I | re.M),
    re.compile(r"^\s*" + _SUBSTANCE_HEAD, re.I | re.M),
    re.compile(r"first information report", re.I),
    re.compile(r"\bfir\s*no\b", re.I),
)
_FIR_SCORE_CAP = 4

# "BETWEEN:" cause title: a petitioner/applicant party block then (within ~15 lines each) an AND line or a respondent block.
_BETWEEN_BLOCK_RE = re.compile(
    r"^[ \t]*BETWEEN[ \t]*:?[ \t]*$(?:\n[^\n]*){0,15}?\n[^\n]*\b(?:PETITIONERS?|APPLICANTS?)\b[^\n]*"
    r"(?:\n[^\n]*){0,15}?\n(?:[ \t]*AND[ \t]*$|[^\n]*\bRESPONDENTS?\b)",
    re.I | re.M)


def _cause_title_window(n_lines: int) -> int:
    """How many leading lines to scan for a cause title. Interpretation of the spec's "min(len//4, 150)":
    scale with document length but never fewer than the old hard 60 (so short docs do not regress) and never
    more than 150. Half the line count is used so a cause title pushed down by a long preamble is still found."""
    return min(max(60, n_lines // 2), 150)


def _fir_score(text: str) -> int:
    """Number of FIR-form markers present in the text (capped)."""
    text = text or ""
    return min(sum(1 for rx in _FIR_FORM_MARKERS if rx.search(text)), _FIR_SCORE_CAP)


def _has_pleading_cause_title(text: str) -> bool:
    """True when the document has a court-pleading cause title ("A versus B" or a BETWEEN: block) near the top."""
    lines = (text or "").splitlines()
    head = lines[:_cause_title_window(len(lines))]
    if any(_VERSUS_LINE_RE.match(l) or _VERSUS_INLINE_RE.match(l) for l in head):
        return True
    return bool(_BETWEEN_BLOCK_RE.search("\n".join(head)))


def _is_fir(text: str) -> bool:
    # A pleading (bail application, writ, appeal) routinely cites an FIR number or quotes FIR text, but is not
    # itself an FIR: require two FIR-form markers AND no pleading cause title.
    return _fir_score(text) >= 2 and not _has_pleading_cause_title(text)


def _fir_block(text: str, heading_rx: str) -> str:
    m = re.search(heading_rx + r"[^\n]*\n(.*?)(?=\n[A-Z][A-Z /&.\-]{6,}[:\n]|\Z)", text, re.S)
    return m.group(1).strip() if m else ""


def _parse_fir(text: str) -> Dict[str, Any]:
    f = {k: "" for k in FIELD_KEYS}
    ps = ""
    m = _FIR_POLICE_STATION_RE.search(text)
    if m:
        ps_line = m.group(1).strip()
        pm = re.match(r"([A-Za-z0-9 ]+?)\s+P\.?\s?S\.?", ps_line)
        ps = (pm.group(1).strip() if pm else "")
    fir = re.search(r"FIR\s*(?:NO\.?|NUMBER)\s*[:.]?\s*([0-9]+\s*/\s*[0-9]{2,4})", text, re.I)
    if fir:
        f["case_number"] = "FIR No. " + re.sub(r"\s+", "", fir.group(1)) + (f" ({ps} P.S.)" if ps else "")
    state = re.search(r"\bState of ([A-Z][a-z]+(?: [A-Z][a-z]+)?)", text)
    if state:
        state_name = "State of " + state.group(1)
    elif m and "," in m.group(1):
        tail = m.group(1).rsplit(",", 1)[1].strip()
        state_name = ("State of " + tail) if re.fullmatch(r"[A-Z][A-Za-z ]{2,30}", tail) else "State"
    else:
        state_name = "State"
    f["petitioner"] = state_name
    names: List[str] = []
    block = _fir_block(text, r"\n(?:" + _ACCUSED_WORDS + r"|SUSPECTED\s*/\s*ACCUSED\s+PERSONS)\s*:?")
    if not block:
        # Numbered government-form layout, e.g. "7. Details of Known / Suspected / Unknown / Accused with full
        # particulars" — bounded by the next top-level numbered item (e.g. "8.") rather than an all-caps heading.
        m = re.search(
            r"\n\s*\d+[.)]\s*Details of Known[^\n]*Accused[^\n]*\n(.*?)(?=\n\s*\d+[.)]\s|\Z)",
            text, re.I | re.S)
        if m:
            block = m.group(1).strip()
    names_by_line: List[str] = []
    for line in block.splitlines():
        am = re.match(r"\s*\(?\d+\)?[.)]?\s*([A-Z][A-Za-z.' ]{2,60}?)\s*(?:\(|,|$)", line)
        if am and not re.match(r"\d+\s+to\s+\d+", line.strip(), re.I):
            nm = am.group(1).strip()
            if nm and not re.match(r"(?:\d|all\b|other|unknown)", nm, re.I):
                names_by_line.append(nm)
    names_inline: List[str] = []
    if block:
        # Numbered entries run together on one line, e.g. "(1) Asit Chakraborty (2) Souparna Roy (3) ..." —
        # split on each "(N)" marker instead of relying on one entry per line.
        for entry in re.split(r"\(\s*\d+\s*\)", block)[1:]:
            nm_m = re.match(r"\s*([A-Z][A-Za-z.' ]{2,60}?)\s*(?:\(|,|@|\bof\b|\bS/o\b|\bD/o\b|\bW/o\b|$)", entry.strip())
            if nm_m:
                nm = nm_m.group(1).strip()
                if nm and not re.match(r"(?:\d|all\b|other|unknown)", nm, re.I):
                    names_inline.append(nm)
    # Whichever extraction actually found more accused names from the text wins — both read the same source,
    # just split on lines vs. on "(N)" markers, and one-per-line docs vs. all-on-one-line docs need different splits.
    names = names_inline if len(names_inline) > len(names_by_line) else names_by_line
    if names:
        f["respondent"] = names[0] + (" & Ors." if len(names) > 1 else "")
        f["case_title"] = f"{state_name} v. {f['respondent']}"
    secs = _fir_block(text, r"\n" + _ACTS_HEAD)
    if secs:
        f["statutory_provisions"] = "; ".join(_clean_line(re.sub(r"^\s*\d+[.)]\s*", "", l)) for l in secs.splitlines() if l.strip())
    if not f["statutory_provisions"]:
        # No "ACTS & SECTIONS:" heading on this form — scan the whole text for "Sections ..."/"u/s ..." mentions instead.
        f["statutory_provisions"] = _statutes(text)
    sub = _fir_block(text, r"\nSUBSTANCE OF COMPLAINT[^\n]*?:?")
    if not sub:
        # No "SUBSTANCE OF COMPLAINT:" heading — fall back to the numbered form's own narrative field.
        sub = _fir_block(text, r"\n\s*\d+[.)]\s*FIR Contents\s*:?")
    f["facts"] = sub
    return f


def _parties(lines: List[str]) -> Tuple[str, str]:
    head = lines[:_cause_title_window(len(lines))]
    for i, line in enumerate(head):
        # A bare "AND" line also separates the party blocks when the cause title opens with "BETWEEN:".
        is_and = bool(re.match(r"^and$", line.strip(), re.I)) and any(
            re.match(r"^between\s*:?$", h.strip(), re.I) for h in head[:i])
        if _VERSUS_LINE_RE.match(line) or is_and:
            before: List[str] = []
            for prev in reversed(head[:i]):
                if re.match(r"^(?:between|in the matter of)\s*:?$", prev.strip(), re.I):
                    break
                if _CASE_NO_RE.search(prev) or _COURT_RE.search(prev) or re.match(r"^(?:petition|writ|appeal|application)\b", prev.strip(), re.I):
                    break
                if prev.lstrip().startswith("(") or len(prev) > 120:
                    break  # parenthetical cause-title note / running prose, not a party name
                before.insert(0, prev)
            after: List[str] = []
            for nxt in head[i + 1:]:
                if not nxt.strip():
                    if after:
                        break
                    continue
                if _heading_kind(nxt) or re.match(r"^(?:writ|petition|application|appeal)\b", nxt.strip(), re.I):
                    break
                after.append(nxt)
                if _PARTY_ROLE_RE.search(nxt):
                    break
            if before and after:
                pet = _party_text(before)
                res = _party_text(after)
                return pet, res
        else:
            m = _VERSUS_INLINE_RE.match(line)
            if m and not _COURT_RE.search(line):
                return _party_text([m.group(1)]), _party_text([m.group(2)])
    return "", ""


def _party_text(chunk: List[str]) -> str:
    names = []
    for l in chunk:
        l = re.sub(r"^(?:between|in the matter of)\s*:?\s*", "", l.strip(), flags=re.I)
        l = _PARTY_TAIL_RE.sub("", l)
        l = _PARTY_ROLE_RE.sub("", l).strip(" .,:;-")
        l = re.sub(r"^\d+[.)]\s*", "", l)
        l = _PERSON_TAIL_RE.sub("", l).strip(" .,:;-")
        if l and not re.match(r"^(?:resident|r/o)\b", l, re.I):
            names.append(l)
    return "; ".join(names[:4]) if len(names) > 1 else (names[0] if names else "")


def _statutes(text: str) -> str:
    found: List[str] = []
    for m in _STATUTE_RE.finditer(text):
        s = _clean_line(m.group(0)).rstrip(",;")
        if s.isupper() or s.casefold() in (x.casefold() for x in found):
            continue
        if s:
            found.append(s)
        if len(found) >= 8:
            break
    return "; ".join(found)


def _parse_generic(text: str) -> Dict[str, Any]:
    """Generic court-pleading field extraction (always run first by extract_fields)."""
    fields: Dict[str, Any] = {k: "" for k in FIELD_KEYS}
    nonempty = [l.strip() for l in text.split("\n") if l.strip()]
    secs = _sections(nonempty)
    for l in nonempty[:25]:
        if _COURT_RE.search(l) and len(l) < 160:
            fields["court"] = re.sub(r"^IN THE\s+", "", l, flags=re.I).strip(" :.-")
            break
    for l in nonempty[:40]:
        if _COURT_RE.search(l):
            continue
        m = _CASE_NO_RE.search(l)
        if m and re.search(r"\d", m.group(0)):
            fields["case_number"] = _clean_line(m.group(0)).strip(" :.-")
            break
    pet, res = _parties(nonempty)
    fields["petitioner"], fields["respondent"] = pet, res
    if pet and res:
        fields["case_title"] = f"{pet} v. {res}"
    fields["facts"] = "\n".join(secs.get("FACTS", [])).strip()
    fields["prayer"] = "\n".join(secs.get("PRAYER", [])).strip()
    fields["petitioner_arguments"] = "\n".join(secs.get("ARGS_P", [])).strip()
    fields["respondent_arguments"] = "\n".join(secs.get("ARGS_R", [])).strip()
    if not fields["facts"]:
        m = re.search(r"(?im)^\s*facts?\s*:\s*(.+)$", text)
        if m:
            fields["facts"] = m.group(1).strip()
    fields["statutory_provisions"] = _statutes(text)
    return fields


def extract_fields(text: str) -> Tuple[Dict[str, str], List[str]]:
    """(fields, issues) read from the (redacted) document text. Anything not present in the text is ''/[]."""
    text = text or ""
    fields = {k: "" for k in FIELD_KEYS}
    if not text.strip():
        return fields, []
    secs = _sections([l.strip() for l in text.split("\n") if l.strip()])

    fields = _parse_generic(text)
    if _is_fir(text):
        # FIR overlay. _is_fir() already guarantees there is no pleading cause title, so on a genuine FIR form the
        # FIR-form readings are authoritative (keeps real-FIR output identical to the pre-classifier behaviour,
        # e.g. the precise ACTS & SECTIONS block beats the whole-text statute scan). The generic pass still
        # contributes every field the FIR pass leaves empty. A misclassified pleading can no longer reach here.
        for k, v in _parse_fir(text).items():
            if v:
                fields[k] = v

    issues: List[str] = []
    for q in _QUESTION_RE.findall(text):
        q = _clean_line(re.sub(r"^\s*\([A-Za-z0-9]+\)\s*", "", q))
        if q and q not in issues:
            issues.append(q)
    if not issues and "ISSUES" in secs:
        for l in secs["ISSUES"]:
            l = _clean_line(re.sub(r"^\s*(?:\(?[A-Za-z0-9]{1,2}[.)]\s+|[-*•]\s+)", "", l))
            if l and l not in issues:
                issues.append(l)
    return fields, issues[:12]


def parse_pleadings(text: str) -> Dict[str, Any]:
    """Legacy /api/upload_pleading shape. Fields not found in the text are empty strings / empty issue text."""
    fields, issues = extract_fields(text)
    return {
        "case_title": fields["case_title"],
        "case_number": fields["case_number"],
        "court": fields["court"],
        "facts": fields["facts"],
        "petitioner_arguments": fields["petitioner_arguments"],
        "respondent_arguments": fields["respondent_arguments"],
        "issues": "\n".join(f"({chr(65 + i)}) {iss}" for i, iss in enumerate(issues)),
        "statutory_provisions": fields["statutory_provisions"],
        "raw_text_length": len(text or ""),
    }


if __name__ == "__main__":
    sample = """
    IN THE HIGH COURT OF JUDICATURE FOR RAJASTHAN AT JODHPUR
    S.B. Civil Writ Petition No. 8920/2023
    Mahaveer Singh v. State of Rajasthan & Ors.

    Whether the impugned suspension order passed without a prior show cause notice violates natural justice?
    """
    print("Parsed Pleading:", parse_pleadings(sample))
